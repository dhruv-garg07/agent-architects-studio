import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation

prs = Presentation(r'C:\Users\Dellg\Downloads\Persistent Context Infrastructure v4.pptx')
print(f"Total slides: {len(prs.slides)}")
print(f"Slide dimensions: {prs.slide_width} x {prs.slide_height}")
print()

for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            texts.append(shape.text.strip())
    
    # Get the main title (usually the first substantial text)
    title = "untitled"
    for t in texts:
        if 10 < len(t) < 80 and t != "Aldenaire International":
            title = t
            break
    
    print(f"Slide {i+1:2d}: {title}")

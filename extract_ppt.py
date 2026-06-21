import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
import json

prs = Presentation(r'C:\Users\Dellg\Downloads\Persistent Context Infrastructure.pptx')

for i, slide in enumerate(prs.slides):
    print(f"\n{'='*60}")
    print(f"SLIDE {i+1}")
    print(f"{'='*60}")
    
    # Get slide layout name
    if slide.slide_layout:
        print(f"Layout: {slide.slide_layout.name}")
    
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            print(f"\n[Shape: {shape.shape_type}, Name: {shape.name}]")
            print(f"Position: left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
            print(f"Text: {shape.text}")

print(f"\n\nTotal slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width}")
print(f"Slide height: {prs.slide_height}")
